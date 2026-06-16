"""
Build the presentation PPTX from the VinUni template, filled with the APEX
content (problem / method / results / demo). Keeps the template's master,
layouts and branding.

Run:  python src/make_pptx.py   ->  outputs/Vietnam-MixedTrafficSim_APEX.pptx
"""

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image

ROOT = Path(__file__).resolve().parent.parent
TEMPLATE = ROOT / "vinuni_template_Presentation1.pptx"
OUT = ROOT / "outputs" / "Vietnam-MixedTrafficSim_APEX.pptx"
# APEX architecture sections rendered from apex_architecture.html
INHERIT_IMG = ROOT / "outputs" / "apex_arch_inherit.png"
CORE_IMG    = ROOT / "outputs" / "apex_arch_core.png"
LEARN_IMG   = ROOT / "outputs" / "apex_arch_learn.png"
DARK = RGBColor(0x0F, 0x14, 0x19)
SHOT = ROOT / "slides" / ("Quang Vu - Vinuni - Robotic - Vietnam Mixed Trafficsim "
                          "- Real time sync driving trip "
                          "[ADWuhLiD450 - 2056x514 - 0m56s].png")

RED = RGBColor(0xB2, 0x1F, 0x1F)
INK = RGBColor(0x22, 0x22, 0x22)
GREEN = RGBColor(0x1E, 0x7A, 0x34)
GREY = RGBColor(0x66, 0x66, 0x66)

# (planner, collisions, <0.3 m fails, min-clr, R_T)
RESULTS = [
    ("IDM (Treiber 2000)",        "5/52",  "7/52",  "5.73 m",  "1.57"),
    ("ORCA/VO (van den Berg 2011)","1/52",  "1/52",  "2.02 m",  "1.34"),
    ("baseline (Frenet, naive)",  "28/52", "30/52", "1.24 m",  "1.27"),
    ("conservative",              "0/52",  "0/52",  "13.96 m", "1.86"),
    ("moto-aware (prev. ours)",   "13/52", "14/52", "2.31 m",  "1.33"),
    ("APEX (predictive, ours)",   "0/52",  "0/52",  "2.16 m",  "1.43"),
]


def layout_by_name(prs, name):
    for lay in prs.slide_layouts:
        if lay.name == name:
            return lay
    return prs.slide_layouts[1]


def image_slide(prs, blank, img):
    """Full-bleed dark slide holding one architecture image, centred."""
    s = prs.slides.add_slide(blank)
    SW, SH = prs.slide_width, prs.slide_height
    rect = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    rect.fill.solid(); rect.fill.fore_color.rgb = DARK; rect.line.fill.background()
    iw, ih = Image.open(img).size
    maxw, maxh = SW - Inches(0.5), SH - Inches(0.5)
    if maxw / (iw / ih) <= maxh:          # width-bound
        pic = s.shapes.add_picture(str(img), 0, 0, width=int(maxw))
    else:                                 # height-bound
        pic = s.shapes.add_picture(str(img), 0, 0, height=int(maxh))
    pic.left = int((SW - pic.width) / 2)
    pic.top = int((SH - pic.height) / 2)
    return s


def set_text(ph, text, size, bold=False, color=INK):
    tf = ph.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color


def bullets(ph, items):
    """items: list of (text, level, color, bold, size)."""
    tf = ph.text_frame
    tf.clear()
    tf.word_wrap = True
    for i, (text, lvl, color, bold, size) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = lvl
        r = p.add_run(); r.text = text
        r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color


def main():
    prs = Presentation(str(TEMPLATE))

    # ── Cover: reuse the template's first slide (3 stacked placeholders) ──
    cover = prs.slides[0]
    phs = sorted(cover.placeholders, key=lambda p: p.top)
    texts = [("Vietnam-MixedTrafficSim", 40, True, RED),
             ("APEX = Anticipatory · Predictive · EXclusion — a safety-constrained "
              "planner for Vietnamese mixed traffic", 18, False, INK),
             ("ELEC5050 Robotics - Group 15   |   Do Minh Phung, Vu Thanh Quang, "
              "Nguyen Quoc Linh", 14, False, GREY)]
    for ph, (t, sz, b, c) in zip(phs, texts):
        set_text(ph, t, sz, b, c)

    tc = layout_by_name(prs, "Title and Content")
    to = layout_by_name(prs, "Title Only")
    blank = layout_by_name(prs, "1_Blank")

    # ── Slide: Problem  (reuse the existing empty Title-and-Content slide) ──
    s = prs.slides[1]
    s.placeholders[0].text = "The Problem: Mixed Traffic in Vietnam"
    bullets(s.placeholders[1], [
        ("AV motion planners assume Western traffic: marked lanes, rule-abiding drivers.", 0, INK, False, 20),
        ("In Vietnam, >90% of vehicles are motorcycles that:", 0, INK, False, 20),
        ("lane-split and squeeze through sub-meter gaps,", 1, GREY, False, 18),
        ("cut in at short range and cross at junctions, ignoring signals.", 1, GREY, False, 18),
        ("Standard planners react too late -> unsafe.", 0, RED, True, 20),
        ("Goal: a planner that ANTICIPATES motorcycle conflicts on the real VinUni road.", 0, GREEN, True, 20),
    ])

    # ── Slide: APEX architecture — inheritance (from apex_architecture.html) ──
    if INHERIT_IMG.exists():
        image_slide(prs, blank, INHERIT_IMG)

    # ── Slide: Real data -> behaviour -> simulation  (from PR2) ──
    s = prs.slides.add_slide(tc)
    s.placeholders[0].text = "Grounded in real Vietnamese data (PR2)"
    bullets(s.placeholders[1], [
        ("5 real car trips around VinUni / Ocean Park (QL5 corridor, roundabout, lakeside, residential, faculty parking).", 0, INK, False, 17),
        ("iPhone 12 Sensor Logger: video + GPS + IMU (100 Hz) + compass + barometer.", 1, GREY, False, 16),
        ("YOLOv8 detects car / motorcycle / pedestrian -> actor counts (density), nearest range (TTC), class mix.", 0, INK, False, 17),
        ("Synchronized 3-panel replay: GPS on the CommonRoad/Lanelet2 map | camera detections | drive stats.", 0, INK, False, 17),
        ("Two demo modes from real density: low-traffic (<=4 actors/frame) vs high-traffic (>4).", 0, INK, False, 17),
        ("Motorcycle calibration: 4 behaviours (lane-splitting, cut-in, close-following, ambiguous priority)", 0, INK, False, 17),
        ("-> lateral velocity, acceleration, heading-change, gap -> stochastic moto agents = APEX's occupancy-ellipse + dynamic safety cost.", 1, GREEN, True, 16),
    ])
    if SHOT.exists():
        s.shapes.add_picture(str(SHOT), Inches(2.8), Inches(5.6), width=Inches(7.7))

    # ── Slide: APEX core (from apex_architecture.html) ──
    if CORE_IMG.exists():
        image_slide(prs, blank, CORE_IMG)

    # ── Slide: Results (table + % improvements) ──
    s = prs.slides.add_slide(to)
    s.placeholders[0].text = "Results: 52 scenarios x 6 planners"
    rows, cols = len(RESULTS) + 1, 5
    tbl = s.shapes.add_table(rows, cols, Inches(0.5), Inches(1.5),
                             Inches(12.3), Inches(3.2)).table
    headers = ["Planner", "Collisions", "<0.3 m fails", "Min-clr", "Mean R_T"]
    for c, h in enumerate(headers):
        cell = tbl.cell(0, c); cell.text = h
        cell.fill.solid(); cell.fill.fore_color.rgb = RED
        p = cell.text_frame.paragraphs[0]; p.runs[0].font.bold = True
        p.runs[0].font.size = Pt(14); p.runs[0].font.color.rgb = RGBColor(255, 255, 255)
    for r, row in enumerate(RESULTS, start=1):
        apex = row[0].startswith("APEX")
        for c, val in enumerate(row):
            cell = tbl.cell(r, c); cell.text = val
            run = cell.text_frame.paragraphs[0].runs[0]
            run.font.size = Pt(13); run.font.bold = apex
            run.font.color.rgb = GREEN if apex else INK
    box = s.shapes.add_textbox(Inches(0.5), Inches(5.0), Inches(12.3), Inches(2.0))
    bullets(box, [
        ("Only conservative & APEX are crash-free - and APEX is the efficient one:", 0, INK, True, 16),
        ("0 collisions / 0 clearance-fails  ->  -100% vs baseline (28) and vs our previous moto-aware (13)", 0, GREEN, False, 15),
        ("23% faster than the only other crash-free planner (R_T 1.43 vs 1.86); beats published IDM & ORCA", 0, GREEN, False, 15),
        ("Metrics (PR2 Sec 6): collision = 0, min clearance > 0.3 m, TTC exposure < 2 s, AEB = 0, efficiency R_T <= 1.25.", 0, GREY, False, 13),
    ])

    # ── Slide: Learning-augmented MPC (next stage, from apex_architecture.html) ──
    if LEARN_IMG.exists():
        image_slide(prs, blank, LEARN_IMG)

    # ── Slide: Demo & next ──
    s = prs.slides.add_slide(tc)
    s.placeholders[0].text = "Live demo & next steps"
    bullets(s.placeholders[1], [
        ("6 scenario videos (4 panels: baseline | ORCA/VO | moto-aware | APEX) - others collide, APEX stays safe:", 0, INK, False, 18),
        ("cut-in, junction crossing, shoulder merge, multi-lane weave, two- and three-motorcycle scenes.", 1, GREY, False, 16),
        ("Aligns with the CommonRoad 2024 competition (winning Frenet paradigm; we add prediction it lacked).", 0, INK, False, 18),
        ("Honest limits (PR2 Sec 7): motorcycle-behaviour calibration in progress; YOLO range is approximate.", 0, RED, False, 16),
        ("Next: calibrate motorcycle behaviour from 200 h moto + 5 car-trip data; CARLA 3D validation.", 0, GREEN, True, 18),
    ])

    OUT.parent.mkdir(exist_ok=True)
    prs.save(str(OUT))
    print(f"Saved {OUT}  ({len(prs.slides)} slides)")


if __name__ == "__main__":
    main()
